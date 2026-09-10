import csv, io, json, re, uuid
from datetime import date, datetime
from openpyxl import load_workbook
from fastapi import APIRouter, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from app.storage import db, get_session, utcnow, log

router = APIRouter()
MAX_FILE_BYTES = 2 * 1024 * 1024

CUSTOMER_COLUMNS = ['company_name','contact_name','phone','email','country','city','activity','service_interest','source','contact_consent','consent_date','notes']
DRIVER_COLUMNS = ['driver_name','whatsapp_phone','vehicle_type','capacity','current_city','preferred_routes','availability','company_name','offer_consent','consent_date','notes']
PROSPECT_COLUMNS = ['company_name','location','sector','lead_status','priority','fit_score','notes']
ALIASES = {
 'اسم الشركة':'company_name','company':'company_name','company name':'company_name','اسم المسؤول':'contact_name','contact':'contact_name','contact name':'contact_name',
 'الجوال':'phone','رقم الجوال':'phone','phone':'phone','البريد':'email','البريد الإلكتروني':'email','email':'email','البلد':'country','الدولة':'country','country':'country','المدينة':'city','city':'city',
 'النشاط':'activity','activity':'activity','الخدمة المطلوبة':'service_interest','service':'service_interest','مصدر البيانات':'source','source':'source',
 'موافقة التواصل':'contact_consent','contact consent':'contact_consent','تاريخ الموافقة':'consent_date','consent date':'consent_date','ملاحظات':'notes','notes':'notes',
 'اسم السائق':'driver_name','driver name':'driver_name','رقم واتساب':'whatsapp_phone','whatsapp':'whatsapp_phone','whatsapp phone':'whatsapp_phone',
 'نوع المركبة':'vehicle_type','vehicle type':'vehicle_type','سعة الحمولة':'capacity','capacity':'capacity','المدينة الحالية':'current_city','current city':'current_city',
 'المسارات المفضلة':'preferred_routes','preferred routes':'preferred_routes','التوفر':'availability','availability':'availability',
 'اسم المؤسسة':'company_name','موافقة استقبال العروض':'offer_consent','offer consent':'offer_consent',
 'الموقع':'location','location':'location','القطاع':'sector','sector':'sector','الحالة':'lead_status','status':'lead_status',
 'الأولوية':'priority','priority':'priority','درجة التوافق':'fit_score','fit score':'fit_score'
}
GCC_PHONE_LENGTHS = {'966':9,'971':9,'973':8,'965':8,'968':8,'974':8}

def _init():
 with db() as c:
  c.execute('''CREATE TABLE IF NOT EXISTS customer_directory(id BIGSERIAL PRIMARY KEY,company_name TEXT NOT NULL,contact_name TEXT,phone TEXT,email TEXT,city TEXT,activity TEXT,service_interest TEXT,source TEXT,contact_consent INTEGER NOT NULL DEFAULT 0,consent_date DATE,notes TEXT,created_at TIMESTAMPTZ NOT NULL,updated_at TIMESTAMPTZ NOT NULL)''')
  c.execute("CREATE UNIQUE INDEX IF NOT EXISTS customer_directory_phone_uq ON customer_directory(phone) WHERE phone IS NOT NULL AND phone<>''")
  c.execute("CREATE UNIQUE INDEX IF NOT EXISTS customer_directory_email_uq ON customer_directory(LOWER(email)) WHERE email IS NOT NULL AND email<>''")
  c.execute('''CREATE TABLE IF NOT EXISTS drivers(id BIGSERIAL PRIMARY KEY,driver_name TEXT NOT NULL,whatsapp_phone TEXT NOT NULL UNIQUE,vehicle_type TEXT NOT NULL,capacity TEXT,current_city TEXT,preferred_routes TEXT,availability TEXT NOT NULL DEFAULT 'متاح',company_name TEXT,offer_consent INTEGER NOT NULL DEFAULT 0,consent_date DATE,notes TEXT,created_at TIMESTAMPTZ NOT NULL,updated_at TIMESTAMPTZ NOT NULL)''')
  c.execute('''CREATE TABLE IF NOT EXISTS data_import_batches(id TEXT PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES users(id),dataset_type TEXT NOT NULL,file_name TEXT NOT NULL,rows_json TEXT NOT NULL,valid_count INTEGER NOT NULL,duplicate_count INTEGER NOT NULL,error_count INTEGER NOT NULL,status TEXT NOT NULL DEFAULT 'preview',created_at TIMESTAMPTZ NOT NULL,committed_at TIMESTAMPTZ)''')
_init()

def esc(v):
 import html
 return html.escape(str(v or ''))

def _session(request): return get_session(request.cookies.get('gla_session'))

def _admin(request):
 s=_session(request)
 if not s:raise HTTPException(401)
 if s.get('role')!='admin':raise HTTPException(403,'Admin role required')
 return s

def _normalize_header(v):
 k=re.sub(r'\s+',' ',str(v or '').strip().lower())
 return ALIASES.get(k,k.replace(' ','_'))

def _phone(v,default_country='966'):
 s=re.sub(r'\D','',str(v or ''))
 if s.startswith('00'):s=s[2:]
 for code,length in GCC_PHONE_LENGTHS.items():
  if s.startswith(code) and len(s)==len(code)+length:return '+'+s
 length=GCC_PHONE_LENGTHS.get(default_country)
 if not length:return ''
 if s.startswith('0') and len(s)==length+1:s=s[1:]
 if len(s)!=length:return ''
 return '+'+default_country+s

def _yes(v): return 1 if str(v or '').strip().lower() in ('1','yes','true','نعم','موافق','موافقة') else 0

def _cell(v):
 if isinstance(v,datetime):return v.date().isoformat()
 if isinstance(v,date):return v.isoformat()
 return str(v or '').strip()

def _valid_date(v):
 if not v:return True
 try:date.fromisoformat(str(v));return True
 except ValueError:return False

def _read_upload(name, raw):
 ext=name.lower().rsplit('.',1)[-1] if '.' in name else ''
 if ext=='csv':
  text=raw.decode('utf-8-sig'); return list(csv.reader(io.StringIO(text)))
 if ext=='xlsx':
  wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);ws=wb.active
  return [[c if c is not None else '' for c in row] for row in ws.iter_rows(values_only=True)]
 raise ValueError('الملف يجب أن يكون XLSX أو CSV')

def _prepare(dataset, matrix, default_country='966'):
 if not matrix:return []
 headers=[_normalize_header(x) for x in matrix[0]]
 allowed=CUSTOMER_COLUMNS if dataset=='customers' else PROSPECT_COLUMNS if dataset=='prospects' else DRIVER_COLUMNS
 out=[];seen=set()
 for n,values in enumerate(matrix[1:],start=2):
  raw={headers[i]:_cell(values[i] if i<len(values) else '') for i in range(len(headers)) if headers[i] in allowed}
  if not any(raw.values()):continue
  if dataset=='prospects':
   errors=[]
   if not raw.get('company_name'):errors.append('اسم الشركة مطلوب')
   key=(raw.get('company_name','').lower(),)
  elif dataset=='customers':
   raw['phone']=_phone(raw.get('phone'),default_country);raw['contact_consent']=_yes(raw.get('contact_consent'))
   errors=[]
   if not raw.get('company_name'):errors.append('اسم الشركة مطلوب')
   if not raw.get('phone') and not raw.get('email'):errors.append('الجوال أو البريد مطلوب')
   if not _valid_date(raw.get('consent_date')):errors.append('التاريخ يجب أن يكون YYYY-MM-DD')
   key=(raw.get('phone') or '',raw.get('email','').lower())
  else:
   raw['whatsapp_phone']=_phone(raw.get('whatsapp_phone'),default_country);raw['offer_consent']=_yes(raw.get('offer_consent'))
   errors=[]
   if not raw.get('driver_name'):errors.append('اسم السائق مطلوب')
   if not raw.get('whatsapp_phone'):errors.append('رقم واتساب غير صالح')
   if not raw.get('vehicle_type'):errors.append('نوع المركبة مطلوب')
   if not _valid_date(raw.get('consent_date')):errors.append('التاريخ يجب أن يكون YYYY-MM-DD')
   key=(raw.get('whatsapp_phone'),)
  duplicate=key in seen and any(key);seen.add(key)
  out.append({'row':n,'data':raw,'errors':errors,'duplicate_in_file':duplicate})
 return out

def page(body):
 return HTMLResponse('''<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>استيراد البيانات</title><style>body{font-family:Arial;background:#07131f;color:#eef6fb;margin:0}.w{max-width:1250px;margin:auto;padding:24px}.card{background:#102536;border:1px solid #28475d;border-radius:16px;padding:20px;margin:14px 0}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:12px}input,select,button{width:100%;padding:11px;border-radius:9px;border:1px solid #36586e;background:#081925;color:#fff}.btn{background:#22c55e;color:#04130a;font-weight:bold;cursor:pointer}.links a{color:#dff6ff;margin-left:12px}table{width:100%;border-collapse:collapse}th,td{padding:9px;border-bottom:1px solid #28475d;text-align:right}.ok{color:#86efac}.bad{color:#fca5a5}.warn{color:#fde68a}.scroll{overflow:auto}</style><div class="w">'''+body+'</div></html>')

@router.get('/data-import',response_class=HTMLResponse)
def import_home(request:Request):
 s=_session(request)
 if not s:return RedirectResponse('/login',303)
 if s.get('role')!='admin':raise HTTPException(403,'Admin role required')
 return page('''<div class="links"><a href="/dashboard">الرئيسية</a><a href="/data-import/template/customers.csv">قالب العملاء CSV</a><a href="/data-import/template/drivers.csv">قالب السائقين CSV</a></div><h1>استيراد قاعدة البيانات</h1><div class="card"><h2>رفع ومعاينة</h2><p>لن تُحفظ البيانات قبل اعتماد المعاينة. الحد الأقصى 2MB. اختر الدولة للأرقام المحلية فقط؛ الأرقام المكتوبة بالمفتاح الدولي تُعتمد تلقائيًا.</p><form method="post" action="/data-import/preview" enctype="multipart/form-data"><div class="grid"><select name="dataset_type"><option value="customers">جهات اتصال العملاء</option><option value="prospects">العملاء المحتملون</option><option value="drivers">السائقون</option></select><select name="default_country"><option value="966">السعودية +966</option><option value="971">الإمارات +971</option><option value="973">البحرين +973</option><option value="965">الكويت +965</option><option value="968">عُمان +968</option><option value="974">قطر +974</option></select><input type="file" name="file" accept=".xlsx,.csv" required></div><p><button class="btn">معاينة الملف</button></p></form></div>''')

@router.get('/data-import/template/{name}')
def template(name:str,request:Request):
 _admin(request)
 if name=='customers.csv':headers=['اسم الشركة','اسم المسؤول','الجوال','البريد الإلكتروني','المدينة','النشاط','الخدمة المطلوبة','مصدر البيانات','موافقة التواصل','تاريخ الموافقة','ملاحظات']
 elif name=='drivers.csv':headers=['اسم السائق','رقم واتساب','نوع المركبة','سعة الحمولة','المدينة الحالية','المسارات المفضلة','التوفر','اسم المؤسسة','موافقة استقبال العروض','تاريخ الموافقة','ملاحظات']
 else:return Response(status_code=404)
 out=io.StringIO();csv.writer(out).writerow(headers)
 return Response('\ufeff'+out.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="{name}"'})

@router.post('/data-import/preview',response_class=HTMLResponse)
async def preview(request:Request,dataset_type:str=Form(...),default_country:str=Form('966'),file:UploadFile=File(...)):
 s=_admin(request)
 if dataset_type not in ('customers','prospects','drivers'):return page('<div class="card bad">نوع البيانات غير صالح.</div>')
 if default_country not in GCC_PHONE_LENGTHS:return page('<div class="card bad">الدولة الافتراضية غير صالحة.</div>')
 raw=await file.read(MAX_FILE_BYTES+1)
 if len(raw)>MAX_FILE_BYTES:return page('<div class="card bad">حجم الملف يتجاوز 2MB.</div>')
 try:items=_prepare(dataset_type,_read_upload(file.filename or '',raw),default_country)
 except Exception as e:return page('<div class="card bad">تعذر قراءة الملف: '+esc(e)+'</div>')
 valid=dupes=errors=0
 with db() as c:
  for x in items:
   d=x['data'];exists=False
   if dataset_type=='drivers' and d.get('whatsapp_phone'):exists=bool(c.execute('SELECT 1 FROM drivers WHERE whatsapp_phone=%s',(d['whatsapp_phone'],)).fetchone())
   if dataset_type=='customers':
    q=[];a=[]
    if d.get('phone'):q.append('phone=%s');a.append(d['phone'])
    if d.get('email'):q.append('LOWER(email)=LOWER(%s)');a.append(d['email'])
    exists=bool(q and c.execute('SELECT 1 FROM customer_directory WHERE '+' OR '.join(q),a).fetchone())
   if dataset_type=='prospects' and d.get('company_name'):
    exists=bool(c.execute('SELECT 1 FROM accounts WHERE LOWER(name)=LOWER(%s)',(d['company_name'],)).fetchone())
   x['duplicate']=exists or x['duplicate_in_file']
   if x['errors']:errors+=1
   elif x['duplicate']:dupes+=1
   else:valid+=1
  bid=str(uuid.uuid4());c.execute('INSERT INTO data_import_batches(id,user_id,dataset_type,file_name,rows_json,valid_count,duplicate_count,error_count,status,created_at) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)',(bid,s['user_id'],dataset_type,file.filename or '',json.dumps(items,ensure_ascii=False),valid,dupes,errors,'preview',utcnow()))
 rows_html=''.join('<tr><td>'+str(x['row'])+'</td><td>'+esc(next(iter(x['data'].values()),''))+'</td><td>'+('<span class="bad">'+esc('، '.join(x['errors']))+'</span>' if x['errors'] else '<span class="warn">مكرر</span>' if x['duplicate'] else '<span class="ok">جاهز</span>')+'</td></tr>' for x in items[:200])
 return page(f'''<h1>معاينة الاستيراد</h1><div class="grid"><div class="card ok">جاهز للحفظ: {valid}</div><div class="card warn">مكرر: {dupes}</div><div class="card bad">أخطاء: {errors}</div></div><div class="card scroll"><table><tr><th>صف Excel</th><th>السجل</th><th>الحالة</th></tr>{rows_html}</table></div><form method="post" action="/data-import/{bid}/commit"><button class="btn" {'disabled' if valid==0 else ''}>اعتماد استيراد السجلات الجاهزة</button></form><p class="links"><a href="/data-import">إلغاء والعودة</a></p>''')

@router.post('/data-import/{batch_id}/commit')
def commit(batch_id:str,request:Request):
 s=_admin(request)
 with db() as c:
  batch=c.execute("SELECT * FROM data_import_batches WHERE id=%s AND user_id=%s AND status='preview' FOR UPDATE",(batch_id,s['user_id'])).fetchone()
  if not batch:return page('<div class="card bad">المعاينة غير موجودة أو تم اعتمادها سابقًا.</div>')
  inserted=0;now=utcnow()
  for x in json.loads(batch['rows_json']):
   if x['errors'] or x.get('duplicate'):continue
   d=x['data']
   if batch['dataset_type']=='drivers':
    saved=c.execute('''INSERT INTO drivers(driver_name,whatsapp_phone,vehicle_type,capacity,current_city,preferred_routes,availability,company_name,offer_consent,consent_date,notes,created_at,updated_at) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,NULLIF(%s,'')::date,%s,%s,%s) ON CONFLICT(whatsapp_phone) DO NOTHING RETURNING id''',(d.get('driver_name'),d.get('whatsapp_phone'),d.get('vehicle_type'),d.get('capacity'),d.get('current_city'),d.get('preferred_routes'),d.get('availability') or 'متاح',d.get('company_name'),d.get('offer_consent',0),d.get('consent_date',''),d.get('notes'),now,now)).fetchone()
   elif batch['dataset_type']=='customers':
    saved=c.execute('''INSERT INTO customer_directory(company_name,contact_name,phone,email,city,activity,service_interest,source,contact_consent,consent_date,notes,created_at,updated_at) VALUES(%s,%s,NULLIF(%s,''),NULLIF(%s,''),%s,%s,%s,%s,%s,NULLIF(%s,'')::date,%s,%s,%s) ON CONFLICT DO NOTHING RETURNING id''',(d.get('company_name'),d.get('contact_name'),d.get('phone',''),d.get('email',''),d.get('city'),d.get('activity'),d.get('service_interest'),d.get('source'),d.get('contact_consent',0),d.get('consent_date',''),d.get('notes'),now,now)).fetchone()
    c.execute('''INSERT INTO accounts(name,country,phone,email,status,notes,created_at,updated_at)
      SELECT %s,%s,NULLIF(%s,''),NULLIF(%s,''),'lead',%s,%s,%s
      WHERE NOT EXISTS (SELECT 1 FROM accounts WHERE LOWER(name)=LOWER(%s))''',
      (d.get('company_name'),d.get('country') or 'السعودية',d.get('phone',''),d.get('email',''),d.get('notes'),now,now,d.get('company_name')))
   else:
    details='القطاع: '+str(d.get('sector') or '')+' | الأولوية: '+str(d.get('priority') or '')+' | درجة التوافق: '+str(d.get('fit_score') or '')
    if d.get('notes'):details+=' | '+str(d.get('notes'))
    saved=c.execute('''INSERT INTO accounts(name,country,domain,status,notes,created_at,updated_at) SELECT %s,%s,%s,%s,%s,%s,%s WHERE NOT EXISTS (SELECT 1 FROM accounts WHERE LOWER(name)=LOWER(%s)) RETURNING id''',(d.get('company_name'),d.get('location'),d.get('sector'),d.get('lead_status') or 'جديد',details,now,now,d.get('company_name'))).fetchone()
   if saved:inserted+=1
  c.execute("UPDATE data_import_batches SET status='committed',committed_at=%s WHERE id=%s",(now,batch_id))
 log(s['user_id'],'data_import',batch['dataset_type'],None,f'Imported {inserted} records from {batch["file_name"]}')
 return page(f'<div class="card ok"><h1>تم الاستيراد</h1><p>تم حفظ {inserted} سجلًا. لم تُحفظ السجلات المكررة أو غير الصالحة.</p><p><a href="/data-import" style="color:white">استيراد ملف آخر</a></p></div>')
