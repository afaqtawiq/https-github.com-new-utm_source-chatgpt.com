"""Pure normalization for a company directory, never an outreach audience."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from urllib.parse import urlsplit

COUNTRIES = {'SA': 'السعودية', 'AE': 'الإمارات', 'QA': 'قطر', 'KW': 'الكويت', 'BH': 'البحرين', 'OM': 'عُمان'}
PHONE_CODES = {'SA': ('966', 9), 'AE': ('971', 9), 'QA': ('974', 8), 'KW': ('965', 8), 'BH': ('973', 8), 'OM': ('968', 8)}
VERIFICATIONS = {'unverified': 'غير متحقق', 'public_contact_verified': 'معلومات عامة موثقة'}
DRIVER_TYPES = {'unreviewed': 'بانتظار تصنيف الإدارة', 'independent_driver': 'سائق مستقل', 'company_driver': 'سائق تابع لشركة', 'carrier_company': 'سجل شركة قديم'}
COLUMNS = ['company_name', 'country', 'city', 'contact_name', 'phone', 'email', 'website', 'preferred_routes', 'vehicle_types', 'capacity', 'source_url', 'maps_url', 'verification_status', 'verified_on', 'notes']
LABELS = dict(zip(COLUMNS, ['اسم الشركة', 'الدولة', 'المدينة', 'جهة الاتصال', 'هاتف الشركة', 'البريد الإلكتروني', 'الموقع الرسمي', 'المسارات والتغطية', 'أنواع المركبات', 'السعة المنشورة', 'رابط المصدر', 'رابط الخرائط', 'حالة توثيق المعلومات العامة', 'تاريخ التحقق', 'ملاحظات المصدر والقيود']))


def normalize_name(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).casefold().split())


def valid_national_number(code, digits, length):
    if code == '971':
        # UAE geographic numbers: one area digit (2/3/4/6/7/9) + 7 subscriber digits.
        # TDRA ICT dictionary / ITU UAE National Numbering Plan; format only, not assignment.
        return ((len(digits) == 8 and digits[0] in '234679') or
                (len(digits) == 9 and digits.startswith('5')))
    return len(digits) == length and not digits.startswith('0')


def normalize_phone(value, country):
    value = str(value or '').strip()
    # CSV export prefixes international numbers to prevent spreadsheet formulas.
    if value.startswith("'+"):
        value = value[1:]
    if not value:
        return ''
    # Do not repair malformed or multiple numbers by guessing.
    if not re.fullmatch(r'\+?[\d\s().-]+', value):
        raise ValueError('هاتف الشركة غير صالح؛ استخدم رقمًا واحدًا')
    digits = ''.join(str(unicodedata.digit(c)) for c in value if c.isdigit())
    international = value.startswith('+') or digits.startswith('00')
    if digits.startswith('00'):
        digits = digits[2:]
    for code, length in PHONE_CODES.values():
        if digits.startswith(code) and valid_national_number(code, digits[len(code):], length):
            return '+' + digits
    code, length = PHONE_CODES[country]
    if not international:
        if digits.startswith('0'):
            digits = digits[1:]
        if valid_national_number(code, digits, length):
            return '+' + code + digits
    raise ValueError('هاتف الشركة غير صالح؛ راجع المصدر ولا تخمن التصحيح')


def valid_url(value):
    if not value:
        return True
    try:
        parsed = urlsplit(value)
        return (parsed.scheme in {'https', 'http'} and bool(parsed.hostname) and
                not parsed.username and not parsed.password and not any(c.isspace() for c in value))
    except ValueError:
        return False


def normalize_record(raw):
    def cell(value):
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value or '').strip()
    record = {key: cell(raw.get(key)) for key in COLUMNS}
    errors = []
    for key, value in record.items():
        if len(value) > (4000 if key == 'notes' else 2000):
            errors.append(LABELS[key] + ': القيمة طويلة جدًا')
    record['country'] = record['country'].upper()
    if not record['company_name']:
        errors.append('اسم الشركة مطلوب')
    if record['country'] not in COUNTRIES:
        errors.append('الدولة يجب أن تكون SA أو AE أو QA أو KW أو BH أو OM')
    else:
        try:
            record['phone'] = normalize_phone(record['phone'], record['country'])
        except ValueError as exc:
            errors.append(str(exc))
    record['email'] = record['email'].lower()
    if record['email'] and not re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', record['email']):
        errors.append('البريد الإلكتروني غير صالح؛ استخدم عنوانًا واحدًا')
    for key in ('website', 'source_url', 'maps_url'):
        if not valid_url(record[key]):
            errors.append(LABELS[key] + ': يلزم رابط http أو https صالح')
    record['verification_status'] = record['verification_status'] or 'unverified'
    if record['verification_status'] not in VERIFICATIONS:
        errors.append('حالة التحقق غير صالحة')
    try:
        if record['verified_on']:
            checked = date.fromisoformat(record['verified_on'])
            if checked > date.today():
                errors.append('تاريخ التحقق لا يمكن أن يكون مستقبليًا')
    except ValueError:
        errors.append('تاريخ التحقق يجب أن يكون YYYY-MM-DD')
    if record['verification_status'] == 'public_contact_verified' and (not record['source_url'] or not record['verified_on']):
        errors.append('توثيق المعلومات العامة يحتاج رابط مصدر وتاريخ تحقق')
    record['normalized_name'] = normalize_name(record['company_name'])
    record['website_domain'] = (urlsplit(record['website']).hostname or '').lower().removeprefix('www.') if valid_url(record['website']) else ''
    return record, errors


def identity_key(record):
    return record['normalized_name'], record['country']


def contact_keys(record):
    return {(key, record[key]) for key in ('phone', 'email', 'website_domain') if record.get(key)}


def prepare_matrix(matrix):
    if not matrix:
        return []
    aliases = {label: key for key, label in LABELS.items()}
    headers = [aliases.get(str(value).strip(), str(value).strip().lower()) for value in matrix[0]]
    if 'company_name' not in headers or 'country' not in headers:
        raise ValueError('استخدم قالب شركات النقل: company_name وcountry مطلوبان')
    if len(set(headers)) != len(headers):
        raise ValueError('الملف يحتوي أعمدة مكررة')
    if len(matrix) > 2001:
        raise ValueError('الحد الأقصى 2000 سجل في الملف')
    items = []
    for number, values in enumerate(matrix[1:], 2):
        raw = {header: values[i] if i < len(values) else '' for i, header in enumerate(headers) if header in COLUMNS}
        if not any(str(v or '').strip() for v in raw.values()):
            continue
        record, errors = normalize_record(raw)
        items.append({'row': number, 'data': record, 'validation_errors': errors})
    return review_candidates(items)


def review_candidates(items, existing_conflicts=None):
    """Reserve contacts only for rows that can really be saved in this preview.

    Recompute after database lookup; an existing duplicate or invalid row must
    not reserve an unused contact and suppress a later unrelated company.
    """
    existing_conflicts = existing_conflicts or {}
    seen, contacts = set(), {}
    for item in items:
        item['errors'] = list(item['validation_errors'])
        item['duplicate'] = False
        if item['errors']:
            continue
        record = item['data']
        key = identity_key(record)
        collision = existing_conflicts.get(item['row'], '')
        if collision == 'duplicate' or key in seen:
            item['duplicate'] = True
            continue
        if collision == 'shared_contact':
            item['errors'].append('جهة اتصال مشتركة مع شركة محفوظة؛ مراجعة يدوية مطلوبة')
            continue
        if any(contact in contacts and contacts[contact] != key for contact in contact_keys(record)):
            item['errors'].append('جهة اتصال مشتركة بين شركتين في الملف؛ راجعها يدويًا')
            continue
        seen.add(key)
        for contact in contact_keys(record):
            contacts[contact] = key
    return items


def csv_text(headers, records):
    def safe(value):
        text = str(value if value is not None else '')
        return "'" + text if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else text
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows([safe(row.get(key, '')) for key in headers] for row in records)
    return '\ufeff' + output.getvalue()


def read_upload(filename, raw):
    extension = filename.lower().rsplit('.',1)[-1]
    if extension == 'csv':
        matrix = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
    elif extension == 'xlsx':
        from zipfile import ZipFile
        from openpyxl import load_workbook
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('ملف Excel كبير بعد فك الضغط')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if (sheet.max_row or 0) > 2001 or (sheet.max_column or 0) > 40:
                raise ValueError('الحد الأقصى 2000 سجل و40 عمودًا')
            matrix = list(sheet.iter_rows(max_row=min(sheet.max_row or 2002,2002),max_col=min(sheet.max_column or 40,40),values_only=True))
        finally:
            workbook.close()
    else:
        raise ValueError('الملف يجب أن يكون CSV أو XLSX')
    if len(matrix)>2001 or any(len(row)>40 for row in matrix):
        raise ValueError('الحد الأقصى 2000 سجل و40 عمودًا')
    return matrix
